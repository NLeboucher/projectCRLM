#!/usr/bin/env python3
"""
Test script for Contracts Management app
This script demonstrates:
1. Reading templates from Contrats_template/ directory
2. Creating a prompt for Claude to extract contract information
"""

import os
from docx import Document

def read_docx_as_text(filepath):
    """Extract text from a .docx file"""
    try:
        doc = Document(filepath)
        text = []
        for paragraph in doc.paragraphs:
            text.append(paragraph.text)
        return '\n'.join(text)
    except Exception as e:
        return f"Error reading file: {e}"

def list_templates():
    """List all templates in Contrats_template/ directory"""
    templates_dir = 'Contrats_template'
    
    if not os.path.exists(templates_dir):
        print(f"Directory {templates_dir} does not exist!")
        return []
    
    templates = []
    print("=== Contract Templates Found ===")
    for filename in os.listdir(templates_dir):
        if filename.endswith('.docx'):
            filepath = os.path.join(templates_dir, filename)
            template_name = os.path.splitext(filename)[0]
            templates.append({
                'name': template_name,
                'filename': filename,
                'filepath': filepath
            })
            print(f"- {template_name}")
    
    return templates

def generate_claude_prompt(template):
    """Generate a prompt for Claude to extract contract information"""
    
    # Read the document content
    content = read_docx_as_text(template['filepath'])
    
    prompt = f"""# Contract Information Extraction Task

You are tasked with analyzing the following contract template and extracting structured information according to the contracttype schema.

## Template Name: {template['name']}

## Template Content (Plain Text):
```
{content}
```

## Your Task:

Based on the contract template above, please extract and structure the following information according to the contracttype.json schema:

1. **Metadata**: Identify the contract type, jurisdiction, governing law
2. **Parties**: Identify placeholders for parties involved (clients, vendors, etc.)
3. **Terms**: Extract duration, renewal terms, notice periods
4. **Financial Terms**: Identify payment terms, schedules, currency
5. **Clauses**: List main contract clauses with their types (scope, payment, confidentiality, liability, etc.)
6. **Obligations**: Extract obligations for each party
7. **Provisions**: Identify termination, confidentiality, liability, dispute resolution clauses
8. **Special Provisions**: Note any IP, force majeure, or other special clauses

Please provide the extracted information in JSON format following the contracttype structure.
"""
    
    return prompt

def main():
    print("=" * 60)
    print("CONTRACTS MANAGEMENT APP - Test & Template Analysis")
    print("=" * 60)
    print()
    
    # List templates
    templates = list_templates()
    print(f"\nTotal templates found: {len(templates)}")
    print()
    
    # Generate Claude prompts for each template
    if templates:
        print("=" * 60)
        print("CLAUDE EXTRACTION PROMPTS")
        print("=" * 60)
        print()
        
        for i, template in enumerate(templates, 1):
            print(f"\n{'=' * 60}")
            print(f"TEMPLATE {i}/{len(templates)}")
            print('=' * 60)
            prompt = generate_claude_prompt(template)
            print(prompt)
            print()

if __name__ == "__main__":
    main()
